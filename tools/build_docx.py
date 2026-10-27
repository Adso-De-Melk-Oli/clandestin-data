"""Construit le modèle Word propre de la carte à partir d'un texte balisé (carte.txt)."""
import sys
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.shared import Pt, Cm, RGBColor

src, out = sys.argv[1], sys.argv[2]

INK = RGBColor(0x1C, 0x1A, 0x17)
BRASS = RGBColor(0x9A, 0x6B, 0x2E)
GREY = RGBColor(0x6B, 0x65, 0x5C)
GREEN = RGBColor(0x2F, 0x6B, 0x3A)
PRICE_TABS = [Cm(11.6), Cm(14.3), Cm(17.0)]  # jusqu'à 3 colonnes de prix, alignées à droite

doc = Document()
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Cm(2)
sec.top_margin = sec.bottom_margin = Cm(1.8)

st = doc.styles
normal = st["Normal"]
normal.font.name = "Georgia"
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(0)


def tune(style, size=None, bold=None, italic=None, color=None, before=0, after=0, indent=None):
    f = style.font
    if size: f.size = Pt(size)
    if bold is not None: f.bold = bold
    if italic is not None: f.italic = italic
    if color is not None: f.color.rgb = color
    f.name = "Georgia"
    pf = style.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    if indent is not None: pf.left_indent = Cm(indent)
    return style


def custom(name, base="Normal", **kw):
    s = st.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    s.base_style = st[base]
    s.quick_style = True
    return tune(s, **kw)


tune(st["Title"], size=26, bold=False, color=INK, before=18, after=4)
tune(st["Heading 1"], size=17, bold=True, color=INK, before=16, after=4)
tune(st["Heading 2"], size=13, bold=True, italic=True, color=BRASS, before=10, after=2)
for h in ("Title", "Heading 1", "Heading 2"):
    st[h].paragraph_format.keep_with_next = True

s_formats = custom("Formats", size=9, bold=True, color=GREY, before=2, after=2)
s_article = custom("Article", size=10.5, bold=True, color=INK, before=6)
s_app = custom("Appellation", size=9.5, bold=True, color=GREY, indent=0.4)
s_cep = custom("Cépages", size=9.5, color=GREY, indent=0.4)
s_desc = custom("Description", size=9.5, italic=True, color=GREY, indent=0.4)
s_comp = custom("Composition", size=9.5, color=GREY, indent=0.9)
s_note = custom("Note", size=9.5, italic=True, color=GREY, after=2)
s_lab = custom("Labels", size=8.5, bold=True, color=GREEN, indent=0.4)

s_article.paragraph_format.keep_with_next = True


def priced(text, values, style):
    """Une tabulation par prix ; taquets posés sur le paragraphe, alignés à droite sur les n dernières colonnes."""
    p = doc.add_paragraph(text + "".join("\t" + v for v in values), style=style)
    for pos in PRICE_TABS[-len(values):]:
        p.paragraph_format.tab_stops.add_tab_stop(pos, WD_TAB_ALIGNMENT.RIGHT)
    return p


fmt_count = 1
for raw in open(src, encoding="utf-8"):
    line = raw.rstrip("\n")
    if not line.strip():
        continue
    tag, _, val = line.partition(": ")
    val = val.strip()
    if tag == "T":
        doc.add_paragraph(val, style="Title")
    elif tag == "H1":
        doc.add_paragraph(val, style="Heading 1")
    elif tag == "H2":
        doc.add_paragraph(val, style="Heading 2")
    elif tag == "F":
        cols = [c.strip() for c in val.split("|")]
        fmt_count = len(cols)
        priced("", cols, "Formats")
    elif tag == "A":
        parts = [p.strip() for p in val.split("|")]
        name, prices, labels = parts[0], parts[1:1 + fmt_count], [p for p in parts[1 + fmt_count:] if p]
        assert len(prices) == fmt_count, line
        priced(name, prices, "Article")
        if labels:
            doc.add_paragraph(", ".join(labels), style="Labels")
    elif tag == "AP":
        doc.add_paragraph(val, style="Appellation")
    elif tag == "CE":
        doc.add_paragraph(val, style="Cépages")
    elif tag == "D":
        doc.add_paragraph(val, style="Description")
    elif tag == "C":
        doc.add_paragraph("– " + val, style="Composition")
    elif tag == "N":
        doc.add_paragraph(val, style="Note")
    else:
        raise SystemExit(f"Balise inconnue : {line}")

doc.core_properties.title = "Carte du Clandestin"
doc.save(out)
print("OK", out)
